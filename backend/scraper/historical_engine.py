import os
import re
import logging
from datetime import date, datetime, timedelta
from typing import List, Dict, Any, Optional
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from sqlalchemy import select, update
from db.database import (
    get_db_sync, HistoricalJob, HistoricalSubJob, HistoricalArticle
)
from scraper.engine import discover_articles, discover_direct_feeds, sanitize_search_keyword
from scraper.search_utils import match_keyword, match_publication_category

logger = logging.getLogger("scraper.historical_engine")

REPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


def check_sub_job_cancellation(db, parent_job_id: str, sub_job_id: str) -> str:
    """
    Checks if either parent job or individual sub-job has been cancelled or paused.
    Returns: 'running', 'paused', or 'cancelled'.
    """
    # Check parent
    res_parent = db.execute(select(HistoricalJob.status).where(HistoricalJob.id == parent_job_id))
    parent_status = res_parent.scalar_one_or_none()
    if parent_status in ("paused", "cancelled"):
        return parent_status

    # Check sub-job
    res_sub = db.execute(select(HistoricalSubJob.status).where(HistoricalSubJob.id == sub_job_id))
    sub_status = res_sub.scalar_one_or_none()
    if sub_status in ("paused", "cancelled"):
        return sub_status

    return "running"


def identify_matched_keywords(text: str, keywords: List[str]) -> List[str]:
    """Identify which target keywords match the article title or metadata."""
    if not text or not keywords:
        return []
    matched = []
    for kw in keywords:
        kw_clean = kw.strip()
        if not kw_clean:
            continue
        if match_keyword(text, kw_clean):
            matched.append(kw_clean)
    return matched


def create_excel_report(file_path: str, title_text: str, articles: List[Dict[str, Any]]) -> str:
    """
    Generates a beautifully styled Excel spreadsheet containing title, publication,
    link, publishing date, and matched keywords.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historical News"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="333333")
    link_font = Font(name="Segoe UI", size=10, color="0563C1", underline="single")
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Header title banner
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = f"HISTORICAL NEWS REPORT — {title_text.upper()}"
    title_cell.font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    headers = ["S.No", "Publishing Date", "Article Title", "Publication", "Matched Keywords", "Article URL"]
    ws.append([]) # Row 2 empty
    ws.append(headers) # Row 3
    ws.row_dimensions[3].height = 24

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num in (1, 2) else "left", vertical="center")

    # Data Rows
    row_idx = 4
    for idx, art in enumerate(articles, 1):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        current_fill = zebra_fill if is_even else None

        row_data = [
            idx,
            art.get("published_date") or "",
            art.get("title") or "",
            art.get("publication") or "Google News",
            art.get("matched_keywords") or "",
            art.get("url") or ""
        ]

        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.font = data_font
            cell.border = thin_border
            if current_fill:
                cell.fill = current_fill

            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6 and val:
                cell.font = link_font
                cell.hyperlink = val

        row_idx += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Cap wide columns
    ws.column_dimensions['C'].width = 50 # Title
    ws.column_dimensions['F'].width = 45 # URL

    wb.save(file_path)
    return file_path


def run_historical_sub_job(parent_job_id: str, sub_job_id: str) -> Dict[str, Any]:
    """
    Executes a metadata-only historical scrape for a single 15-day sub-window.
    """
    with get_db_sync() as db:
        # Load Parent and SubJob
        res_parent = db.execute(select(HistoricalJob).where(HistoricalJob.id == parent_job_id))
        parent_job = res_parent.scalar_one_or_none()

        res_sub = db.execute(select(HistoricalSubJob).where(HistoricalSubJob.id == sub_job_id))
        sub_job = res_sub.scalar_one_or_none()

        if not parent_job or not sub_job:
            logger.error(f"Historical job/sub-job not found: {parent_job_id} / {sub_job_id}")
            return {"status": "failed", "error": "Job not found"}

        # Update sub-job to running
        db.execute(update(HistoricalSubJob).where(HistoricalSubJob.id == sub_job_id).values(
            status="running", started_at=datetime.now()
        ))
        db.commit()

        keywords = [k.strip() for k in parent_job.keywords.split(",") if k.strip()]
        date_from = sub_job.date_from
        date_to = sub_job.date_to
        cumulative_urls = set()
        sub_articles = []

        curr_date = date_from
        while curr_date <= date_to:
            status_check = check_sub_job_cancellation(db, parent_job_id, sub_job_id)
            if status_check in ("paused", "cancelled"):
                logger.info(f"Sub-job {sub_job_id} stopped via status '{status_check}'.")
                db.execute(update(HistoricalSubJob).where(HistoricalSubJob.id == sub_job_id).values(status=status_check))
                db.commit()
                return {"status": status_check, "count": len(sub_articles)}

            # Discover items via Google News RSS & direct feeds metadata
            discovered = discover_articles(keywords, curr_date, "IN", "india", f"hist_{sub_job_id}", cumulative_urls, is_brand_track=False, sector=parent_job.name)
            direct = []
            if curr_date >= date.today():
                direct = discover_direct_feeds(keywords, curr_date, f"hist_{sub_job_id}", cumulative_urls, sector=parent_job.name)

            all_raw = discovered + direct

            for item in all_raw:
                url = item.get("url") or ""
                title = item.get("title") or ""
                if not url or not title:
                    continue

                # Identify matched keywords
                matched = identify_matched_keywords(title, keywords)
                matched_str = ", ".join(matched) if matched else parent_job.name

                agency = item.get("agency") or item.get("source") or "Google News"
                pub_date_str = curr_date.strftime("%Y-%m-%d")

                sub_articles.append({
                    "title": title,
                    "publication": agency,
                    "url": url,
                    "published_date": pub_date_str,
                    "matched_keywords": matched_str
                })

            curr_date += timedelta(days=1)

        # Save extracted articles into database
        for art in sub_articles:
            db.execute(
                HistoricalArticle.__table__.insert().values(
                    parent_job_id=parent_job_id,
                    sub_job_id=sub_job_id,
                    title=art["title"],
                    publication=art["publication"],
                    url=art["url"],
                    published_date=art["published_date"],
                    matched_keywords=art["matched_keywords"],
                    created_at=datetime.now()
                )
            )

        # Generate 15-Day Sub-Window Excel File
        sub_excel_filename = f"historical_{parent_job_id}_window_{sub_job.window_index}.xlsx"
        sub_excel_path = os.path.join(REPORTS_DIR, sub_excel_filename)
        title_text = f"{parent_job.name} ({date_from} to {date_to})"
        create_excel_report(sub_excel_path, title_text, sub_articles)

        # Update SubJob state
        db.execute(update(HistoricalSubJob).where(HistoricalSubJob.id == sub_job_id).values(
            status="completed",
            articles_found=len(sub_articles),
            excel_file_path=sub_excel_path,
            completed_at=datetime.now()
        ))

        # Re-query total completed articles & sub-jobs for Parent Job
        res_completed_count = db.execute(
            select(HistoricalSubJob).where(HistoricalSubJob.parent_job_id == parent_job_id).where(HistoricalSubJob.status == "completed")
        )
        completed_subs = res_completed_count.scalars().all()
        completed_count = len(completed_subs)

        # Fetch all historical articles for cumulative master report
        res_all_arts = db.execute(
            select(HistoricalArticle).where(HistoricalArticle.parent_job_id == parent_job_id).order_by(HistoricalArticle.published_date.desc())
        )
        all_job_articles = [
            {
                "title": a.title,
                "publication": a.publication,
                "url": a.url,
                "published_date": a.published_date,
                "matched_keywords": a.matched_keywords
            } for a in res_all_arts.scalars().all()
        ]

        # Generate Cumulative Master Excel File
        master_excel_filename = f"historical_{parent_job_id}_MASTER.xlsx"
        master_excel_path = os.path.join(REPORTS_DIR, master_excel_filename)
        master_title = f"{parent_job.name} MASTER ({parent_job.date_from} to {parent_job.date_to})"
        create_excel_report(master_excel_path, master_title, all_job_articles)

        # Check if all sub-jobs are finished
        all_subs_res = db.execute(select(HistoricalSubJob).where(HistoricalSubJob.parent_job_id == parent_job_id))
        all_sub_jobs = all_subs_res.scalars().all()

        all_done = all(sj.status in ("completed", "cancelled", "paused") for sj in all_sub_jobs)
        parent_status = "completed" if (all_done and any(sj.status == "completed" for sj in all_sub_jobs)) else parent_job.status

        db.execute(update(HistoricalJob).where(HistoricalJob.id == parent_job_id).values(
            completed_sub_jobs=completed_count,
            total_articles=len(all_job_articles),
            master_excel_path=master_excel_path,
            status=parent_status,
            completed_at=datetime.now() if all_done else None
        ))
        db.commit()

        return {
            "status": "completed",
            "count": len(sub_articles),
            "excel_path": sub_excel_path,
            "master_path": master_excel_path
        }
